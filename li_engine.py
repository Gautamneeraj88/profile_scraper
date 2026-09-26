"""
Engine for the LinkedIn Enricher.
=================================

Headless, GUI-free and print-free: everything here can be driven equally by the
desktop app or the command line.  Progress is reported by *returning* values and
by `logging.getLogger("enrich")`, never by printing, and the row loop is a
generator the caller pumps so a UI can show per-candidate progress, pause, and
stop.

Three things in here are load-bearing and were learned the hard way:

* `settle()` scrolls `main#workspace`, not the document.  LinkedIn lazy-loads its
  sections inside an inner container, so window-level scrolling loads nothing.
* Sections are found by their heading text and entries by `/school/` and
  `/company/` anchor hrefs, because LinkedIn's class names are obfuscated and
  rotate between deploys.
* All writable state lives under `app_dir()`, never beside the script.  Under a
  frozen PyInstaller build the script directory is a temp folder that is wiped on
  exit, which would destroy the saved browser session after every single run.
"""

from __future__ import annotations

import contextlib
import csv
import io
import json
import logging
import os
import random
import re
import sys
import threading
import time
import traceback
import urllib.parse
import urllib.request
import uuid
from dataclasses import dataclass, field, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterator

import li_fields as LF
from li_fields import Ctx, FIELDS, Overwrite

log = logging.getLogger("enrich")

ENGINE_VERSION = "2.0.0"
APP_VENDOR = "GradNext"
APP_NAME = "LinkedInEnricher"

USER_AGENT = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")

PAGE_TIMEOUT = 45_000
SETTLE_MAX_ROUNDS = 40
LOGIN_WAIT_SECONDS = 15 * 60
PAGE_RETRIES = 2

BLANK_TOKENS = {"", "-", "--", "n/a", "na", "#n/a", "#na", "none", "null", "nil",
                "tbd", "?", "not available", "#ref!", "#value!"}


# ===========================================================================
# section 1 -- errors
# ===========================================================================
class EngineError(Exception):
    """Anything the user should be shown rather than a traceback."""


class InputError(EngineError):
    """The spreadsheet or URL could not be used."""


class SheetAccessError(InputError):
    """A Google Sheet that is not readable."""


class MappingError(EngineError):
    """The column mapping is unusable."""


class OutputLockedError(EngineError):
    """The output file is open in Excel."""


class Cancelled(Exception):
    """Raised inside the engine when the user presses Stop."""


# ===========================================================================
# section 2 -- paths.  Must not depend on where the script lives.
# ===========================================================================
def app_dir() -> Path:
    """Per-user writable directory for config, cache, logs and the browser profile.

    A frozen build unpacks to a temp directory that is deleted on exit, so
    anything stored next to the executable -- crucially the Chrome profile
    holding the LinkedIn session -- would not survive a single run.
    """
    override = os.environ.get("LI_ENRICHER_HOME")
    if override:
        base = Path(override).expanduser()
    elif sys.platform == "win32":
        base = Path(os.environ.get("APPDATA") or Path.home() / "AppData" / "Roaming")
        base = base / APP_VENDOR / APP_NAME
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support" / APP_VENDOR / APP_NAME
    else:
        base = Path(os.environ.get("XDG_CONFIG_HOME") or Path.home() / ".config")
        base = base / APP_VENDOR / APP_NAME
    base.mkdir(parents=True, exist_ok=True)
    return base


def config_path() -> Path:
    return app_dir() / "config.json"


def profile_dir() -> Path:
    p = app_dir() / "chrome-profile"
    p.mkdir(parents=True, exist_ok=True)
    return p


def cache_dir() -> Path:
    p = app_dir() / "cache" / f"profiles-v{LF.RECORD_SCHEMA}"
    p.mkdir(parents=True, exist_ok=True)
    return p


def log_dir() -> Path:
    p = app_dir() / "logs"
    p.mkdir(parents=True, exist_ok=True)
    return p


def tmp_dir() -> Path:
    p = app_dir() / "tmp"
    p.mkdir(parents=True, exist_ok=True)
    return p


def default_output_dir() -> Path:
    docs = Path.home() / "Documents"
    return (docs if docs.is_dir() else Path.home()) / APP_VENDOR


def setup_logging(verbose: bool = False, extra: logging.Handler | None = None) -> Path:
    """Attach a fresh per-run log file.  Idempotent: our own handlers are replaced
    rather than stacked, so a second run in the same process does not duplicate
    every line or leak file handles."""
    for handler in list(log.handlers):
        if getattr(handler, "_enrich_owned", False):
            log.removeHandler(handler)
            with contextlib.suppress(Exception):
                handler.close()
    path = log_dir() / f"run-{datetime.now():%Y%m%d-%H%M%S}.log"
    file_handler = logging.FileHandler(path, encoding="utf-8")
    file_handler.setFormatter(
        logging.Formatter("%(asctime)s %(levelname)-7s %(message)s"))
    file_handler.setLevel(logging.DEBUG)
    file_handler._enrich_owned = True
    log.addHandler(file_handler)
    if verbose:
        stream = logging.StreamHandler()
        stream.setFormatter(logging.Formatter("  . %(message)s"))
        stream.setLevel(logging.INFO)
        stream._enrich_owned = True
        log.addHandler(stream)
    if extra is not None:
        extra._enrich_owned = True
        log.addHandler(extra)
    log.setLevel(logging.DEBUG)
    log.propagate = False
    return path


class RedactFilter(logging.Filter):
    """Keep the password out of the log, even if Playwright echoes a fill() call."""

    def __init__(self) -> None:
        super().__init__()
        self._secrets: list[str] = []

    def add(self, secret: str) -> None:
        if secret and len(secret) >= 4:
            self._secrets.append(secret)

    def filter(self, record: logging.LogRecord) -> bool:
        if self._secrets:
            try:
                message = record.getMessage()
                for secret in self._secrets:
                    if secret in message:
                        record.msg = message.replace(secret, "***")
                        record.args = ()
            except Exception:
                pass
        return True


REDACT = RedactFilter()
log.addFilter(REDACT)


# ===========================================================================
# section 3 -- run control: pause, stop, cancellable waits
# ===========================================================================
class Control:
    """Cooperative pause/stop shared between the UI thread and the engine thread.

    Built on threading.Event only, so the command line can use it for Ctrl-C and
    the GUI can drive it without the engine importing Qt.
    """

    def __init__(self) -> None:
        self._stop = threading.Event()
        self._go = threading.Event()
        self._go.set()

    def stop(self) -> None:
        self._stop.set()
        self._go.set()          # wake anything parked in a pause

    def pause(self) -> None:
        self._go.clear()

    def resume(self) -> None:
        self._go.set()

    def reset(self) -> None:
        self._stop.clear()
        self._go.set()

    @property
    def stopping(self) -> bool:
        return self._stop.is_set()

    @property
    def paused(self) -> bool:
        return not self._go.is_set()

    def sleep(self, seconds: float, tick: float = 0.2) -> None:
        """Wait `seconds` of *running* time.  Paused time does not count down.

        Raises Cancelled promptly on stop, so pressing Stop during a 15-second
        pacing wait takes effect within ~200 ms instead of blocking the run.
        """
        remaining = max(0.0, float(seconds))
        while True:
            if self._stop.is_set():
                raise Cancelled()
            if not self._go.is_set():
                self._go.wait(tick)
                continue
            if remaining <= 0:
                return
            slice_ = min(tick, remaining)
            if self._stop.wait(slice_):
                raise Cancelled()
            remaining -= slice_

    def checkpoint(self) -> None:
        """Stop/pause boundary with no delay."""
        self.sleep(0.0)


# ===========================================================================
# section 4 -- configuration
# ===========================================================================
CONFIG_VERSION = 1


@dataclass
class ColumnMap:
    header: str
    field: str | None = None
    enabled: bool = True

    def to_json(self) -> dict:
        return {"header": self.header, "field": self.field, "enabled": self.enabled}

    @staticmethod
    def from_json(raw: dict) -> "ColumnMap":
        key = raw.get("field")
        if key is not None and key not in FIELDS:
            log.warning("config refers to unknown field %r -- treated as a manual column", key)
            key = None
        return ColumnMap(header=str(raw.get("header") or "").strip(),
                         field=key, enabled=bool(raw.get("enabled", True)))


@dataclass
class AppConfig:
    input_mode: str = "sheet"                 # sheet | file | urls
    sheet_url: str = ""
    file_path: str = ""
    urls: list[str] = field(default_factory=list)

    username: str = ""
    remember_password: bool = True

    min_delay: float = 8.0
    max_delay: float = 15.0
    long_pause_every: int = 25
    long_pause: tuple[float, float] = (45.0, 90.0)

    headless: bool = True
    name_search: bool = True
    use_cache: bool = True
    cache_ttl_days: int = 14
    limit: int = 0
    max_relogins: int = 1

    output_folder: str = ""
    output_stem: str = "candidates_enriched"
    output_xlsx: bool = True
    output_csv: bool = True
    highlight_needs_review: bool = True

    columns: list[ColumnMap] = field(default_factory=list)
    hidden_columns: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        if not self.columns:
            self.columns = default_columns()
        if not self.output_folder:
            self.output_folder = str(default_output_dir())

    # -- derived ----------------------------------------------------------
    @property
    def enabled_fields(self) -> list[str]:
        return [c.field for c in self.columns if c.enabled and c.field]

    @property
    def needs(self) -> frozenset:
        return LF.required_visits(self.enabled_fields)

    def output_headers(self) -> list[str]:
        return [c.header for c in self.columns if c.enabled]

    def validate(self) -> list[str]:
        problems = []
        headers = [c.header.strip() for c in self.columns if c.enabled]
        if not headers:
            problems.append("No columns are enabled -- there would be nothing to write.")
        if any(not h for h in headers):
            problems.append("Every column needs a header name.")
        lowered = [h.lower() for h in headers]
        dupes = sorted({h for h in lowered if lowered.count(h) > 1})
        if dupes:
            problems.append("These column headers are used more than once: "
                            + ", ".join(dupes))
        if self.min_delay > self.max_delay:
            problems.append("The shortest delay cannot be longer than the longest delay.")
        if self.min_delay < 0:
            problems.append("Delays cannot be negative.")
        if self.input_mode == "sheet" and not self.sheet_url.strip():
            problems.append("No Google Sheet link has been entered.")
        if self.input_mode == "file" and not self.file_path.strip():
            problems.append("No Excel or CSV file has been chosen.")
        if self.input_mode == "urls" and not [u for u in self.urls if u.strip()]:
            problems.append("No LinkedIn profile links have been pasted in.")
        if not self.output_xlsx and not self.output_csv:
            problems.append("Choose at least one output format (Excel or CSV).")
        return problems

    def warnings(self) -> list[str]:
        """Advisory notes -- shown to the user, but they do not block a run."""
        notes = []
        if self.min_delay < 4:
            notes.append(
                f"A gap of only {self.min_delay:g}s between profiles is risky: LinkedIn "
                f"may pause your account. 8 seconds or more is much safer.")
        pages = len(LF.required_visits(self.enabled_fields))
        if pages > 5:
            notes.append(
                f"The fields you have chosen need {pages} LinkedIn pages per candidate. "
                f"The more pages per person, the likelier LinkedIn is to step in.")
        experimental = [FIELDS[f].label for f in self.enabled_fields
                        if FIELDS[f].experimental]
        if experimental:
            notes.append("These fields have not been verified against real profiles "
                         "yet, so check them: " + ", ".join(experimental))
        return notes

    # -- persistence ------------------------------------------------------
    def to_json(self) -> dict:
        data = asdict(self)
        data["columns"] = [c.to_json() for c in self.columns]
        data["long_pause"] = list(self.long_pause)
        data["config_version"] = CONFIG_VERSION
        data["engine_version"] = ENGINE_VERSION
        return data

    @staticmethod
    def from_json(raw: dict) -> "AppConfig":
        raw = dict(raw or {})
        version = int(raw.pop("config_version", CONFIG_VERSION) or CONFIG_VERSION)
        if version > CONFIG_VERSION:
            raise EngineError(
                f"This settings file was written by a newer version of the app "
                f"(format {version}, this app understands {CONFIG_VERSION}). "
                f"Please update the app, or use Reset to defaults.")
        raw.pop("engine_version", None)
        columns = [ColumnMap.from_json(c) for c in (raw.pop("columns", None) or [])]
        lp = raw.pop("long_pause", None)
        known = {f for f in AppConfig.__dataclass_fields__}
        unknown = set(raw) - known
        for key in unknown:
            log.info("ignoring unknown setting %r", key)
        cfg = AppConfig(**{k: v for k, v in raw.items() if k in known})
        if columns:
            cfg.columns = columns
        if isinstance(lp, (list, tuple)) and len(lp) == 2:
            cfg.long_pause = (float(lp[0]), float(lp[1]))
        return cfg

    def save(self, path: Path | None = None) -> Path:
        path = path or config_path()
        tmp = path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(self.to_json(), indent=2), encoding="utf-8")
        os.replace(tmp, path)
        log.info("settings saved to %s", path)
        return path

    @staticmethod
    def load(path: Path | None = None) -> "AppConfig":
        path = path or config_path()
        if not path.exists():
            return AppConfig()
        try:
            return AppConfig.from_json(json.loads(path.read_text(encoding="utf-8")))
        except EngineError:
            raise
        except Exception as exc:
            raise EngineError(f"Could not read the settings file {path}: {exc}") from exc


def default_columns() -> list[ColumnMap]:
    """The nine GradNext columns plus the meta columns: today's CLI output exactly."""
    return [ColumnMap(header=FIELDS[key].label, field=key, enabled=True)
            for key in LF.DEFAULT_PRESET]


# ===========================================================================
# section 5 -- credential storage
# ===========================================================================
KEYRING_SERVICE = f"{APP_VENDOR} {APP_NAME}"


@dataclass
class Credentials:
    username: str = ""
    password: str = ""

    @property
    def usable(self) -> bool:
        return bool(self.username.strip() and self.password)


class SecretStore:
    """Saves the LinkedIn password in the operating system's keystore.

    Windows Credential Manager, macOS Keychain or the Linux Secret Service --
    whichever `keyring` finds.  When there is no keystore we deliberately do NOT
    invent our own: a password 'encrypted' with a key sitting next to it is
    obfuscation, not protection, and the honest fallback is to ask the user to
    sign in by hand.
    """

    def __init__(self) -> None:
        self._backend = None
        self.reason = ""
        try:
            import keyring
            from keyring.backends import fail as _fail
            backend = keyring.get_keyring()
            if isinstance(backend, _fail.Keyring):
                self.reason = "no system keystore is available on this computer"
            else:
                self._backend = keyring
        except Exception as exc:                       # pragma: no cover
            self.reason = f"the keyring library is unavailable ({exc})"

    @property
    def available(self) -> bool:
        return self._backend is not None

    @property
    def name(self) -> str:
        if not self.available:
            return "Not saved"
        cls = type(self._backend.get_keyring()).__module__
        if "Windows" in cls:
            return "Windows Credential Manager"
        if "macOS" in cls:
            return "macOS Keychain"
        return "System keystore"

    def save(self, username: str, password: str) -> bool:
        if not self.available or not username or not password:
            return False
        try:
            self._backend.set_password(KEYRING_SERVICE, username, password)
            return True
        except Exception as exc:
            log.warning("could not save the password: %s", exc)
            return False

    def load(self, username: str) -> str:
        if not self.available or not username:
            return ""
        try:
            return self._backend.get_password(KEYRING_SERVICE, username) or ""
        except Exception as exc:
            log.warning("could not read the saved password: %s", exc)
            return ""

    def delete(self, username: str) -> None:
        if not self.available or not username:
            return
        with contextlib.suppress(Exception):
            self._backend.delete_password(KEYRING_SERVICE, username)


# ===========================================================================
# section 6 -- page extraction JavaScript
#
# Everything here keys off heading text and anchor hrefs, never class names:
# LinkedIn's classes are obfuscated and rotate between deploys, while
# `a[href*="/school/"]` cannot be obfuscated away.  The `href|lines` dedupe key
# also suppresses LinkedIn's aria-hidden duplicated text, so do not switch these
# to textContent.
# ===========================================================================
PROFILE_EXTRACT = r"""(sectionSpecs) => {
  const clean = (s) => (s || "").split("\n").map((x) => x.trim()).filter(Boolean);

  const secByHeading = (src) => {
    const re = new RegExp(src, "i");
    return Array.from(document.querySelectorAll("main section")).find((s) => {
      const h = s.querySelector("h2");
      return h && re.test((h.innerText || "").trim());
    }) || null;
  };

  const anchorEntries = (sec, sel) => {
    const out = [];
    const seen = new Set();
    for (const a of sec.querySelectorAll(sel)) {
      if (!a.innerText.trim()) continue;
      const lines = clean(a.innerText).slice(0, 10);
      const key = a.href + "|" + lines.join("|");
      if (seen.has(key)) continue;
      seen.add(key);
      out.push({ lines, link: a.getAttribute("href") });
    }
    return out;
  };

  // Sections such as Skills and Languages have no /school/ or /company/ anchor,
  // so fall back to the list items and drop the aria-hidden duplicate lines.
  const listEntries = (sec) => {
    const out = [];
    const seen = new Set();
    for (const li of sec.querySelectorAll("li")) {
      const lines = clean(li.innerText).filter((l, i, arr) => i === 0 || l !== arr[i - 1]);
      if (!lines.length) continue;
      const key = lines.join("|");
      if (seen.has(key) || key.length < 2) continue;
      // skip wrapper <li>s that contain other <li>s
      if (li.querySelector("li")) continue;
      seen.add(key);
      out.push({ lines: lines.slice(0, 6), link: null });
    }
    return out;
  };

  const showAll = (sec) => {
    const a = Array.from(sec.querySelectorAll("a")).find((x) =>
      /show all|see all/i.test(x.innerText || "") || /\/details\//.test(x.href || ""));
    return a ? a.href : null;
  };

  const sections = {};
  for (const spec of sectionSpecs) {
    const sec = secByHeading(spec.heading);
    if (!sec) { sections[spec.name] = null; continue; }
    let entries;
    if (spec.anchor) entries = anchorEntries(sec, spec.anchor);
    else entries = listEntries(sec);
    if (!entries.length && spec.anchor) entries = listEntries(sec);
    sections[spec.name] = {
      entries: entries,
      show_all: showAll(sec),
      heading: (sec.querySelector("h2") || {}).innerText || "",
    };
  }

  const topSec = document.querySelector("main section");
  const topLines = topSec ? clean(topSec.innerText) : [];

  // Location and headline: try the design-system utility classes first, then let
  // the Python side fall back to reading the top-card lines around "Contact info".
  const locEl = document.querySelector(
    "main section .text-body-small.inline.t-black--light.break-words, " +
    "main section span.text-body-small:not(.inline)");
  const headEl = document.querySelector("main section .text-body-medium.break-words");
  const photo = document.querySelector(
    "main img.pv-top-card-profile-picture__image, " +
    "main img.pv-top-card-profile-picture__image--show, " +
    "main button img[width='200'], main section img[width='200']");

  return {
    name: topLines[0] || "",
    topcard: topLines.slice(0, 16),
    headline: headEl ? (headEl.innerText || "").trim() : "",
    location: locEl ? (locEl.innerText || "").trim() : "",
    photo: photo ? (photo.getAttribute("src") || "") : "",
    open_to_work: !!document.querySelector(
      "main .pv-open-to-carousel-card, main [data-test-id*='open-to'], " +
      "main img[alt*='OPEN_TO_WORK'], main .pv-member-badge--for-top-card"),
    about: (() => {
      const s = secByHeading("^about");
      return s ? clean(s.innerText).slice(1).join(" ").slice(0, 1200) : "";
    })(),
    headings: Array.from(document.querySelectorAll("main section h2")).map((h) =>
      h.innerText.trim()),
    sections: sections,
    authwall: /join linkedin|sign in to view/i.test(document.body.innerText.slice(0, 1500)),
  };
}"""

DETAIL_EXTRACT = r"""(spec) => {
  const clean = (s) => (s || "").split("\n").map((x) => x.trim()).filter(Boolean);
  const main = document.querySelector("main");
  if (!main) return null;
  const seen = new Set();
  const out = [];
  if (spec.anchor) {
    for (const a of main.querySelectorAll(spec.anchor)) {
      if (!a.innerText.trim()) continue;
      const lines = clean(a.innerText).slice(0, 10);
      const key = a.href + "|" + lines.join("|");
      if (seen.has(key)) continue;
      seen.add(key);
      out.push({ lines, link: a.getAttribute("href") });
    }
  }
  if (!out.length) {
    for (const li of main.querySelectorAll("li")) {
      if (li.querySelector("li")) continue;
      const lines = clean(li.innerText).filter((l, i, arr) => i === 0 || l !== arr[i - 1]);
      if (!lines.length) continue;
      const key = lines.join("|");
      if (seen.has(key) || key.length < 2) continue;
      seen.add(key);
      out.push({ lines: lines.slice(0, 6), link: null });
    }
  }
  return { entries: out, url: location.href };
}"""

CONTACT_EXTRACT = r"""() => {
  const clean = (s) => (s || "").split("\n").map((x) => x.trim()).filter(Boolean);
  const dlg = document.querySelector('[role="dialog"]') ||
              document.querySelector(".artdeco-modal") ||
              document.querySelector("section.pv-contact-info");
  if (!dlg) return { found: false };

  const out = { found: true, email: "", phone: "", websites: [], twitter: "" };
  const text = clean(dlg.innerText).join("\n");
  out.text = text.slice(0, 1500);

  const mail = dlg.querySelector('a[href^="mailto:"]');
  if (mail) out.email = decodeURIComponent((mail.getAttribute("href") || "").slice(7))
                          .split("?")[0].trim();
  const tel = dlg.querySelector('a[href^="tel:"]');
  if (tel) out.phone = decodeURIComponent((tel.getAttribute("href") || "").slice(4)).trim();

  for (const sec of dlg.querySelectorAll("section, li")) {
    const h = sec.querySelector("h3, .pv-contact-info__header");
    const label = h ? h.innerText.trim().toLowerCase() : "";
    if (!label) continue;
    const values = clean(sec.innerText).filter((l) => l.toLowerCase() !== label);
    if (/phone/.test(label) && values.length && !out.phone) {
      out.phone = values[0].replace(/\s*\((home|mobile|work|cell)\)\s*/i, "").trim();
    }
    if (/website/.test(label)) {
      for (const a of sec.querySelectorAll("a[href]")) {
        const href = a.getAttribute("href");
        if (href && !/linkedin\.com/.test(href)) out.websites.push(href);
      }
    }
    if (/twitter|^x$/.test(label) && values.length && !out.twitter) out.twitter = values[0];
  }

  if (!out.email) {
    const m = text.match(/[\w.+-]+@[\w-]+\.[\w.-]{2,}/);
    if (m) out.email = m[0];
  }
  if (!out.phone && /phone/i.test(text)) {
    const m = text.match(/(\+?\d[\d\s().-]{6,17}\d)/);
    if (m) out.phone = m[1].trim();
  }
  return out;
}"""

SEARCH_EXTRACT = r"""() => {
  const clean = (s) => (s || "").split("\n").map((x) => x.trim()).filter(Boolean);
  const main = document.querySelector("main") || document.body;
  const seen = new Set();
  const out = [];
  for (const a of main.querySelectorAll('a[href*="/in/"]')) {
    const href = (a.getAttribute("href") || "").split("?")[0];
    const m = href.match(/\/in\/([^/?#]+)/);
    if (!m) continue;
    const slug = m[1];
    if (seen.has(slug)) continue;
    let box = a;
    for (let i = 0; i < 5 && box.parentElement; i++) {
      box = box.parentElement;
      if (box.innerText && box.innerText.trim().split("\n").length >= 3) break;
    }
    const lines = clean(box.innerText).slice(0, 8);
    if (!lines.length) continue;
    seen.add(slug);
    out.push({ slug, lines });
  }
  return {
    results: out.slice(0, 8),
    blocked: /unusual activity|verify you.?re a human|try again later/i.test(
      document.body.innerText.slice(0, 2000)),
  };
}"""

# Which anchor selector identifies real entries in each section.
SECTION_ANCHOR = {
    LF.VISIT_EDUCATION: 'a[href*="/school/"]',
    LF.VISIT_EXPERIENCE: 'a[href*="/company/"]',
}


def section_specs(visits) -> list[dict]:
    """The section descriptors handed to the page-extraction JavaScript."""
    specs = []
    for visit in LF.VISIT_ORDER:
        if visit in (LF.VISIT_PROFILE, LF.VISIT_CONTACT):
            continue
        heading = LF.SECTION_HEADING.get(visit)
        if not heading:
            continue
        # Always read education and experience previews -- they are the common case
        # and cost nothing extra once the page is open.
        if visit not in visits and visit not in (LF.VISIT_EDUCATION, LF.VISIT_EXPERIENCE):
            continue
        specs.append({"name": visit, "heading": heading,
                      "anchor": SECTION_ANCHOR.get(visit)})
    return specs


# ===========================================================================
# section 7 -- browser plumbing
# ===========================================================================
FIND_SCROLLER = """() => {
  const cands = [document.querySelector('main#workspace'), document.querySelector('main'),
                 document.scrollingElement];
  for (const el of cands) {
    if (el && el.scrollHeight > el.clientHeight + 50) { window.__sc = el; return el.tagName; }
  }
  window.__sc = document.scrollingElement;
  return 'fallback';
}"""

SCROLL_STEP = """() => {
  const el = window.__sc || document.scrollingElement;
  el.scrollTop = Math.min(el.scrollTop + Math.round(el.clientHeight * 0.75), el.scrollHeight);
  const m = document.querySelector('main');
  return {top: el.scrollTop, sh: el.scrollHeight, ch: el.clientHeight, len: m ? m.innerText.length : 0};
}"""


def settle(page, control: Control, max_rounds: int = SETTLE_MAX_ROUNDS) -> None:
    """Scroll LinkedIn's *inner* container until lazy-loaded sections stop growing.

    LinkedIn scrolls `main#workspace`, not the document, so window-level scrolling
    does nothing at all.  This is the single thing that makes the scrape work.
    """
    page.wait_for_timeout(random.randint(1500, 2400))
    page.evaluate(FIND_SCROLLER)
    last_len, last_sh, stable = -1, -1, 0
    for _ in range(max_rounds):
        control.checkpoint()          # Stop aborts within one round, not 40
        r = page.evaluate(SCROLL_STEP)
        page.wait_for_timeout(random.randint(450, 800))
        at_end = r["top"] + r["ch"] >= r["sh"] - 40
        if r["len"] == last_len and r["sh"] == last_sh and at_end:
            stable += 1
            if stable >= 3:
                break
        else:
            stable = 0
        last_len, last_sh = r["len"], r["sh"]
    page.evaluate("(window.__sc || document.scrollingElement).scrollTop = 0")
    page.wait_for_timeout(random.randint(400, 800))


def clean_stale_locks(directory: Path) -> None:
    """Remove Chrome's lock files left behind by a killed run.

    Only when no live process owns them: evicting a running sibling browser would
    corrupt the shared profile.
    """
    owner = directory / ".owner"
    if owner.exists():
        try:
            pid = int(owner.read_text().strip() or 0)
        except (ValueError, OSError):
            pid = 0
        if pid and pid != os.getpid() and _pid_alive(pid):
            raise EngineError(
                "The app already has a browser open (another window, or a run still "
                "finishing). Please close it and try again.")
    for name in ("SingletonLock", "SingletonCookie", "SingletonSocket",
                 "RunningChromeVersion"):
        target = directory / name
        try:
            if target.is_symlink() or target.exists():
                target.unlink()
                log.info("removed stale lock %s", name)
        except OSError as exc:
            log.warning("could not remove %s: %s", name, exc)
    with contextlib.suppress(OSError):
        owner.write_text(str(os.getpid()))


def _pid_alive(pid: int) -> bool:
    if sys.platform == "win32":                        # pragma: no cover
        import ctypes
        handle = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)
        if handle:
            ctypes.windll.kernel32.CloseHandle(handle)
            return True
        return False
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    return True


def find_chrome_binary() -> str | None:
    """Prefer the full Chrome-for-Testing build over the headless shell."""
    override = os.environ.get("PLAYWRIGHT_BROWSERS_PATH")
    roots = [Path(override)] if override else [
        Path.home() / "Library" / "Caches" / "ms-playwright",
        Path.home() / ".cache" / "ms-playwright",
        Path(os.environ.get("LOCALAPPDATA", "")) / "ms-playwright",
    ]
    patterns = [
        "chromium-*/chrome-mac-arm64/Google Chrome for Testing.app/Contents/MacOS/Google Chrome for Testing",
        "chromium-*/chrome-mac/Google Chrome for Testing.app/Contents/MacOS/Google Chrome for Testing",
        "chromium-*/chrome-linux/chrome",
        "chromium-*/chrome-win/chrome.exe",
    ]
    best = None
    for root in roots:
        if not root.is_dir():
            continue
        for pattern in patterns:
            for hit in sorted(root.glob(pattern)):
                if best is None or str(hit) > str(best):
                    best = hit
    return str(best) if best else None


def browser_installed() -> bool:
    return find_chrome_binary() is not None


def launch_context(pw, headless: bool):
    directory = profile_dir()
    clean_stale_locks(directory)
    kwargs = dict(
        user_data_dir=str(directory),
        headless=headless,
        user_agent=USER_AGENT,
        viewport={"width": 1440, "height": 900},
        locale="en-US",
        args=["--disable-blink-features=AutomationControlled"],
    )
    binary = find_chrome_binary()
    if binary:
        try:
            return pw.chromium.launch_persistent_context(executable_path=binary, **kwargs)
        except Exception as exc:
            log.warning("bundled Chrome failed (%s), falling back", exc)
    try:
        return pw.chromium.launch_persistent_context(channel="chrome", **kwargs)
    except Exception:
        return pw.chromium.launch_persistent_context(**kwargs)


BLOCK_URL_PATTERNS = (
    ("authwall", re.compile(r"/authwall|/uas/login", re.I)),
    ("checkpoint", re.compile(r"/checkpoint", re.I)),
)


def is_blocked(page) -> str | None:
    url = page.url or ""
    for label, pattern in BLOCK_URL_PATTERNS:
        if pattern.search(url):
            return label
    if re.search(r"linkedin\.com/login", url, re.I):
        return "authwall"
    try:
        head = page.evaluate("() => document.body ? document.body.innerText.slice(0,1500) : ''")
    except Exception:
        return None
    if re.search(r"unusual activity|verify you.?re a human|quick security check", head, re.I):
        return "captcha"
    if re.search(r"join linkedin|sign in to view|you.?ve reached the (weekly|monthly) limit",
                 head, re.I):
        return "authwall"
    return None


def goto(page, url: str, control: Control, retries: int = PAGE_RETRIES) -> None:
    last = None
    for attempt in range(retries + 1):
        control.checkpoint()
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT)
            return
        except Exception as exc:
            last = exc
            log.warning("could not open %s (try %s of %s): %s",
                        url, attempt + 1, retries + 1, exc)
            if attempt < retries:
                control.sleep(5 + 10 * attempt)
    raise last


NAV_SELECTORS = ("#global-nav", "nav[aria-label*='Primary']", ".global-nav",
                 "[data-test-global-nav]", "input[placeholder*='Search']")


def session_is_live(page, control: Control) -> bool:
    """Actually load the feed.  A stale li_at cookie looks valid but redirects to
    the login page, so checking the cookie alone gives false positives."""
    try:
        goto(page, "https://www.linkedin.com/feed/", control)
    except Cancelled:
        raise
    except Exception as exc:
        log.warning("could not reach the LinkedIn feed: %s", exc)
        return False
    if is_blocked(page):
        log.info("session check: blocked at %s", page.url)
        return False
    for selector in NAV_SELECTORS:
        try:
            if page.query_selector(selector):
                return True
        except Exception:
            continue
    log.info("session check: no signed-in navigation at %s", page.url)
    return False


def has_cookie(ctx) -> bool:
    try:
        cookies = ctx.cookies("https://www.linkedin.com")
    except Exception:
        return False
    return any(c.get("name") == "li_at" and c.get("value") for c in cookies)


# ===========================================================================
# section 8 -- signing in
#
# The window is ALWAYS visible for a sign-in.  LinkedIn challenges automated
# logins often, so the app fills the form and then waits for the human to clear
# whatever LinkedIn asks for.  Silently submitting credentials headlessly fails
# frequently and raises the chance of the account being restricted.
# ===========================================================================
LOGIN_URL = "https://www.linkedin.com/login"
SEL_USERNAME = "input#username, input[name='session_key']"
SEL_PASSWORD = "input#password, input[name='session_password']"
SEL_SUBMIT = "button[type=submit][data-litms-control-urn], button[type=submit], .btn__primary--large"
SEL_ERRORS = ("#error-for-password", "#error-for-username", "div[error-for]",
              ".form__label--error", ".alert-content")
SEL_CHALLENGE = ("#input__phone_verification_pin", "input[name='pin']",
                 "#two-step-challenge", "iframe[title*='challenge']",
                 "iframe[src*='captcha']", "#captcha-internal")
CHALLENGE_URL_RE = re.compile(r"/checkpoint/(challenge|lg|challengesV2)", re.I)

LOGIN_HINTS = {
    "typed": "Signing you in ...",
    "manual": "Please sign in to LinkedIn in the browser window that just opened.",
    "otp": "LinkedIn has sent you a verification code. Please type it into the "
           "browser window -- this app will carry on by itself once you are in.",
    "challenge": "LinkedIn wants to check it is really you. Please finish that in "
                 "the browser window; the app is waiting.",
    "captcha": "LinkedIn is showing a puzzle to prove you are human. Please solve "
               "it in the browser window; the app is waiting.",
}


@dataclass
class LoginEvent:
    stage: str
    message: str


def classify_login_page(page) -> str:
    url = page.url or ""
    if re.search(r"/feed|/mynetwork|/in/", url, re.I):
        for selector in NAV_SELECTORS:
            with contextlib.suppress(Exception):
                if page.query_selector(selector):
                    return "success"
    if CHALLENGE_URL_RE.search(url):
        for selector in SEL_CHALLENGE:
            with contextlib.suppress(Exception):
                if page.query_selector(selector):
                    if "captcha" in selector:
                        return "captcha"
                    return "otp" if "pin" in selector else "challenge"
        return "challenge"
    for selector in SEL_ERRORS:
        with contextlib.suppress(Exception):
            node = page.query_selector(selector)
            if node and (node.inner_text() or "").strip():
                return "bad_credentials"
    with contextlib.suppress(Exception):
        if page.query_selector(SEL_USERNAME):
            return "form"
    return "unknown"


def read_login_error(page) -> str:
    for selector in SEL_ERRORS:
        with contextlib.suppress(Exception):
            node = page.query_selector(selector)
            if node:
                text = (node.inner_text() or "").strip()
                if text:
                    return text
    return "LinkedIn did not accept those sign-in details."


def assisted_login(pw, creds: Credentials | None, control: Control,
                   emit: Callable[[LoginEvent], None]) -> bool:
    """Open a visible window, pre-fill the credentials if we have them, then wait
    for the human to complete anything LinkedIn asks for."""
    ctx = launch_context(pw, headless=False)
    try:
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        with contextlib.suppress(Exception):
            page.bring_to_front()
        goto(page, LOGIN_URL, control)

        if has_cookie(ctx) and classify_login_page(page) == "success":
            emit(LoginEvent("success", "Already signed in."))
            return True

        if creds and creds.usable:
            REDACT.add(creds.password)
            try:
                page.wait_for_selector(SEL_USERNAME, timeout=15_000)
                page.fill(SEL_USERNAME, creds.username)
                page.click(SEL_PASSWORD)
                page.type(SEL_PASSWORD, creds.password, delay=45)
                emit(LoginEvent("typed", LOGIN_HINTS["typed"]))
                page.click(SEL_SUBMIT)
            except Exception as exc:
                log.warning("could not fill the sign-in form automatically: %s", exc)
                emit(LoginEvent("manual", LOGIN_HINTS["manual"]))
        else:
            emit(LoginEvent("manual", LOGIN_HINTS["manual"]))

        deadline = time.monotonic() + LOGIN_WAIT_SECONDS
        announced: set[str] = set()
        while time.monotonic() < deadline:
            control.sleep(1.0)            # Stop aborts here within a second
            stage = classify_login_page(page)
            if stage == "success" and has_cookie(ctx):
                emit(LoginEvent("success", "Signed in. Thank you."))
                page.wait_for_timeout(2500)
                return True
            if stage == "bad_credentials":
                emit(LoginEvent("bad_credentials", read_login_error(page)))
                return False
            if stage in ("otp", "challenge", "captcha") and stage not in announced:
                announced.add(stage)
                emit(LoginEvent(stage, LOGIN_HINTS[stage]))
        emit(LoginEvent("timeout", "Timed out waiting for the sign-in to finish."))
        return False
    finally:
        with contextlib.suppress(Exception):
            ctx.close()


class Session:
    """Owns the browser context so it can be torn down and rebuilt for a re-login."""

    def __init__(self, pw, headless: bool, control: Control,
                 creds: Credentials | None = None, max_relogins: int = 1):
        self.pw = pw
        self.headless = headless
        self.control = control
        self.creds = creds
        self.max_relogins = max_relogins
        self.ctx = None
        self.page = None
        self.relogins = 0

    def open(self) -> None:
        self.ctx = launch_context(self.pw, self.headless)
        self.page = self.ctx.pages[0] if self.ctx.pages else self.ctx.new_page()
        self.page.set_default_timeout(PAGE_TIMEOUT)

    def close(self) -> None:
        if self.ctx is not None:
            with contextlib.suppress(Exception):
                self.ctx.close()
            self.ctx, self.page = None, None
        with contextlib.suppress(OSError):
            (profile_dir() / ".owner").unlink()

    def live(self) -> bool:
        return self.page is not None and session_is_live(self.page, self.control)

    def relogin(self, emit: Callable[[LoginEvent], None]) -> bool:
        if self.relogins >= self.max_relogins:
            return False
        self.relogins += 1
        log.info("session expired mid-run; asking the user to sign in again")
        self.close()
        ok = assisted_login(self.pw, self.creds, self.control, emit)
        self.open()
        return bool(ok and self.live())


# ===========================================================================
# section 9 -- the record cache
#
# Caches the RAW scraped record, not derived values.  That is what makes a
# user-configurable field mapping cheap: adding or renaming a column re-derives
# offline in milliseconds with no LinkedIn traffic at all.  A cached record also
# means a run interrupted by a block resumes for free.
# ===========================================================================
class RecordCache:
    def __init__(self, ttl_days: int = 14, enabled: bool = True):
        self.ttl = max(0, int(ttl_days)) * 86400
        self.enabled = enabled

    def _path(self, slug: str) -> Path:
        safe = re.sub(r"[^A-Za-z0-9_.-]", "_", slug)[:120]
        return cache_dir() / f"{safe}.json"

    def get(self, slug: str, needs: frozenset) -> tuple[dict | None, frozenset]:
        """Return (record, visits_still_needed).

        A record fetched when only Education was mapped has no contact block, so
        it must NOT satisfy a run that now wants Email.  Rather than throwing it
        away we top it up: reuse what is there and open only the missing pages.
        """
        if not self.enabled:
            return None, needs
        path = self._path(slug)
        if not path.exists():
            return None, needs
        try:
            envelope = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None, needs
        if int(envelope.get("schema", 0)) != LF.RECORD_SCHEMA:
            log.info("cache for %s is from an older version -- refetching", slug)
            return None, needs
        if self.ttl and (time.time() - float(envelope.get("fetched_at", 0))) > self.ttl:
            log.info("cache for %s has expired -- refetching", slug)
            return None, needs
        record = envelope.get("record")
        if not isinstance(record, dict) or not record.get("ok"):
            return None, needs
        have = frozenset(envelope.get("visits") or ())
        missing = frozenset(needs) - have
        record["from_cache"] = not missing
        record["cached_visits"] = sorted(have)
        return record, missing

    def put(self, slug: str, visits: frozenset, record: dict, pageviews: int) -> None:
        if not self.enabled or not record.get("ok"):
            return
        envelope = {"schema": LF.RECORD_SCHEMA, "engine": ENGINE_VERSION,
                    "fetched_at": time.time(), "visits": sorted(visits),
                    "pageviews": pageviews, "record": record}
        path = self._path(slug)
        tmp = path.with_suffix(".json.tmp")
        try:
            tmp.write_text(json.dumps(envelope, indent=1, ensure_ascii=False),
                           encoding="utf-8")
            os.replace(tmp, path)
        except OSError as exc:
            log.warning("could not cache %s: %s", slug, exc)

    def clear(self) -> int:
        removed = 0
        for path in cache_dir().glob("*.json"):
            with contextlib.suppress(OSError):
                path.unlink()
                removed += 1
        return removed


# ===========================================================================
# section 10 -- scraping one profile, opening only the pages that are needed
# ===========================================================================
SLUG_RE = re.compile(r"linkedin\.com/(?:[a-z]{2,3}/)?(?:in|pub)/([^/?#\s]+)", re.I)
BARE_SLUG_RE = re.compile(r"^[a-z0-9](?:[a-z0-9-]{1,98}[a-z0-9])?$", re.I)
BAD_SLUGS = {"in", "pub", "company", "school", "feed", "jobs", "mynetwork", "me",
             "unavailable", "login", "signup"}


def is_blank(value) -> bool:
    if value is None:
        return True
    return str(value).strip().lower() in BLANK_TOKENS


def extract_slug(value) -> str | None:
    """Pull the /in/<slug> handle out of whatever shape the cell is in."""
    if is_blank(value):
        return None
    raw = str(value).strip().strip("<>").replace("​", "")
    m = SLUG_RE.search(raw)
    if m:
        slug = m.group(1)
    else:
        if any(ch in raw for ch in "/ .@") or "linkedin" in raw.lower():
            return None
        slug = raw
    slug = urllib.parse.unquote(slug).strip().strip("/")
    if not slug or len(slug) > 100 or slug.lower() in BAD_SLUGS:
        return None
    if not BARE_SLUG_RE.match(slug):
        return None
    return slug


def profile_url(slug: str) -> str:
    return f"https://www.linkedin.com/in/{slug}/"


def scrape_profile(page, slug: str, needs: frozenset, control: Control,
                   cache: RecordCache, page_gap: tuple[float, float] = (1.5, 3.0)) -> dict:
    """Fetch one profile, opening only the pages the mapped fields require.

    `page_gap` is the pause between the extra pages of a single profile; it comes
    from the user's pacing setting so a deliberately slow run stays slow
    throughout, and a faster one is not silently held back here.
    """
    cached, missing = cache.get(slug, needs)
    if cached is not None and not missing:
        log.info("%s served from cache", slug)
        return cached

    record = cached or {"slug": slug, "url": profile_url(slug), "ok": False,
                        "blocked": None}
    record.setdefault("slug", slug)
    record.setdefault("url", profile_url(slug))
    to_visit = missing if cached is not None else frozenset(needs)
    pageviews = 0

    if cached is None or LF.VISIT_PROFILE in to_visit or "profile" not in record:
        goto(page, profile_url(slug), control)
        pageviews += 1
        blocked = is_blocked(page)
        if blocked:
            record["blocked"] = blocked
            return record
        settle(page, control)
        record["profile"] = page.evaluate(PROFILE_EXTRACT, section_specs(needs))
        record["final_url"] = page.url
        if (record["profile"] or {}).get("authwall"):
            record["blocked"] = "authwall"
            return record
        # hoist the section previews to the top level for the field context
        for name, section in ((record["profile"].get("sections") or {}).items()):
            if section:
                record.setdefault(name, {}).update(section)

    ctx = Ctx(record)

    # Education and experience: open the details page only when the field set asks
    # for it AND LinkedIn actually truncated the list behind a "Show all" link.
    for visit in (LF.VISIT_EDUCATION, LF.VISIT_EXPERIENCE):
        if visit not in needs:
            continue
        section = record.get(visit) or {}
        if not section.get("show_all"):
            continue
        try:
            goto(page, section["show_all"], control)
            pageviews += 1
            if is_blocked(page):
                record["blocked"] = is_blocked(page)
                return record
            settle(page, control)
            full = page.evaluate(DETAIL_EXTRACT,
                                 {"anchor": SECTION_ANCHOR.get(visit)})
            if full and full.get("entries"):
                record[visit]["entries"] = full["entries"]
            control.sleep(random.uniform(*page_gap))
        except Cancelled:
            raise
        except Exception as exc:
            log.warning("details page failed for %s/%s: %s", slug, visit, exc)

    # The other sections each cost a page, so skip anyone whose profile did not
    # show that heading at all.
    for visit in LF.VISIT_ORDER:
        if visit in (LF.VISIT_PROFILE, LF.VISIT_CONTACT, LF.VISIT_EDUCATION,
                     LF.VISIT_EXPERIENCE) or visit not in needs:
            continue
        existing = record.get(visit) or {}
        if not ctx.has_section(visit):
            log.info("%s has no %s section -- skipping that page", slug, visit)
            record.setdefault(visit, {"entries": [], "show_all": None})
            continue
        if existing.get("entries") and not existing.get("show_all"):
            continue        # the preview was already complete
        try:
            url = existing.get("show_all") or (
                f"https://www.linkedin.com/in/{slug}/details/{LF.SECTION_URL[visit]}/")
            goto(page, url, control)
            pageviews += 1
            if is_blocked(page):
                record["blocked"] = is_blocked(page)
                return record
            settle(page, control)
            full = page.evaluate(DETAIL_EXTRACT, {"anchor": None})
            if full:
                record[visit] = {"entries": full.get("entries") or [], "show_all": None}
            control.sleep(random.uniform(*page_gap))
        except Cancelled:
            raise
        except Exception as exc:
            log.warning("%s page failed for %s: %s", visit, slug, exc)

    if LF.VISIT_CONTACT in needs and (cached is None or LF.VISIT_CONTACT in to_visit):
        try:
            goto(page, f"https://www.linkedin.com/in/{slug}/overlay/contact-info/", control)
            pageviews += 1
            page.wait_for_timeout(random.randint(1800, 2800))
            if is_blocked(page):
                record["blocked"] = is_blocked(page)
                return record
            record["contact"] = page.evaluate(CONTACT_EXTRACT)
        except Cancelled:
            raise
        except Exception as exc:
            log.warning("contact info failed for %s: %s", slug, exc)
            record["contact"] = {"found": False}

    record["ok"] = True
    record["pageviews"] = pageviews
    record["from_cache"] = False
    visits_have = frozenset(record.get("cached_visits") or ()) | frozenset(needs)
    cache.put(slug, visits_have, record, pageviews)
    return record


# ===========================================================================
# section 11 -- finding a profile from a name
# ===========================================================================
STOP_TOKENS = {"mr", "mrs", "ms", "dr", "prof", "the", "jr", "sr", "ii", "iii"}


def name_tokens(name: str) -> list[str]:
    toks = re.split(r"[^a-z]+", (name or "").lower())
    return [t for t in toks if len(t) > 1 and t not in STOP_TOKENS]


def email_tokens(email: str) -> list[str]:
    local = (email or "").split("@")[0]
    local = re.sub(r"\d+", " ", local)
    return [t for t in re.split(r"[^a-z]+", local.lower()) if len(t) > 2]


def score_match(result: dict, name: str, email: str = "") -> float:
    lines = result.get("lines") or []
    result_name = lines[0] if lines else ""
    blob = " ".join(lines).lower()
    want, got = name_tokens(name), name_tokens(result_name)
    if not want or not got:
        return 0.0

    score = 0.0
    if set(want) == set(got):
        score += 0.60
    elif all(t in got for t in want):
        score += 0.45
    elif len(want) > 1 and want[-1] in got:
        score += 0.20
    elif want[0] in got:
        score += 0.15
    else:
        return 0.0

    hints = email_tokens(email)
    if hints and any(h in got or h in blob for h in hints):
        score += 0.25
    if re.search(r"\b(1st|2nd|3rd)\b", blob):
        score += 0.10
    if re.search(r"\bindia\b", blob):
        score += 0.05

    score = min(score, 0.99)
    # a one-word query could match thousands of people; never call that confident
    if len(want) == 1:
        score = min(score, 0.70)
    return score


def band(score: float) -> str:
    if score >= 0.80:
        return "high"
    if score >= 0.55:
        return "medium"
    return "low"


def search_people(page, name: str, control: Control) -> dict:
    url = ("https://www.linkedin.com/search/results/people/?keywords="
           + urllib.parse.quote(name))
    goto(page, url, control)
    blocked = is_blocked(page)
    if blocked:
        return {"blocked": blocked, "results": []}
    settle(page, control, max_rounds=12)
    data = page.evaluate(SEARCH_EXTRACT) or {}
    if data.get("blocked"):
        return {"blocked": "captcha", "results": []}
    return {"blocked": None, "results": data.get("results") or []}


def resolve_missing_profile(page, name: str, email: str, control: Control) -> dict:
    """Best-effort name search.  LinkedIn cannot be searched by email or phone, so
    this is a guess by name and is always reported with a confidence band."""
    found = search_people(page, name, control)
    if found["blocked"]:
        return {"slug": None, "confidence": "", "blocked": found["blocked"],
                "note": f"search blocked ({found['blocked']})"}
    scored = sorted(((score_match(r, name, email), r) for r in found["results"]),
                    key=lambda pair: pair[0], reverse=True)
    scored = [(s, r) for s, r in scored if s > 0]
    if not scored:
        return {"slug": None, "confidence": "", "blocked": None,
                "note": "no LinkedIn profile matched this name"}

    best_score, best = scored[0]
    confidence = band(best_score)
    note = ""
    if len(scored) > 1 and best_score - scored[1][0] < 0.10:
        confidence = {"high": "medium", "medium": "low"}.get(confidence, "low")
        note = "several people share this name"
    return {"slug": best["slug"], "confidence": confidence, "blocked": None,
            "note": note, "matched_name": (best.get("lines") or [""])[0],
            "score": round(best_score, 2)}


# ===========================================================================
# section 12 -- reading the input
# ===========================================================================
GSHEET_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([a-zA-Z0-9-_]+)")


def gsheet_export_url(url: str) -> str:
    m = GSHEET_RE.search(url or "")
    if not m:
        raise InputError("That does not look like a Google Sheets link.")
    gid = re.search(r"[#&?]gid=(\d+)", url)
    return (f"https://docs.google.com/spreadsheets/d/{m.group(1)}"
            f"/export?format=csv&gid={gid.group(1) if gid else '0'}")


def download_gsheet(url: str) -> Path:
    export = gsheet_export_url(url)
    log.info("downloading the sheet from %s", export)
    request = urllib.request.Request(export, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            body = response.read()
            final_url = response.geturl()
    except Exception as exc:
        raise SheetAccessError(
            "Could not download that Google Sheet.\n\n"
            f"{exc}\n\n"
            "Check the link, and make sure the sheet is shared as "
            "'Anyone with the link can view'. You can also download it as Excel "
            "(File > Download > Microsoft Excel) and use the file instead.") from exc
    if b"<html" in body[:400].lower() or "accounts.google.com" in final_url:
        raise SheetAccessError(
            "That Google Sheet is not readable without signing in.\n\n"
            "Either share it as 'Anyone with the link can view', or download it as "
            "Excel (File > Download > Microsoft Excel) and choose that file instead.")
    # a unique name per run, so two runs never overwrite each other's input
    destination = tmp_dir() / f"sheet-{uuid.uuid4().hex[:8]}.csv"
    destination.write_bytes(body)
    return destination


def load_rows(path: Path) -> tuple[list[dict], list[str]]:
    """Read a csv/tsv/xlsx into plain string rows plus the header order."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:                     # pragma: no cover
            raise InputError("Reading Excel files needs the openpyxl library.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        headers: list[str] = []
        for raw in rows_iter:
            if raw and any(c is not None and str(c).strip() for c in raw):
                headers = ["" if c is None else str(c).strip() for c in raw]
                break
        rows = []
        for raw in rows_iter:
            if raw is None or not any(c is not None and str(c).strip() for c in raw):
                continue
            record = {}
            for i, header in enumerate(headers):
                if not header:
                    continue
                value = raw[i] if i < len(raw) else None
                record[header] = _cell_text(value)
            rows.append(record)
        workbook.close()
        return rows, [h for h in headers if h]

    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    delimiter = "\t" if suffix == ".tsv" else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    headers = [h.strip() for h in (reader.fieldnames or []) if h and h.strip()]
    rows = []
    for record in reader:
        clean = {(k or "").strip(): _cell_text(v)
                 for k, v in record.items() if k and k.strip()}
        if any(clean.values()):
            rows.append(clean)
    return rows, headers


def _cell_text(value) -> str:
    """Excel stores phone numbers as floats: 917874835898.0 -> 917874835898."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def rows_from_urls(urls: list[str]) -> tuple[list[dict], list[str]]:
    """Build a fresh table from a pasted list of profile links.

    Deliberately strict: only text that really contains linkedin.com/in/ counts.
    A bare word is a valid slug inside a spreadsheet's LinkedIn column, but when
    someone pastes a block of prose here we must not invent profiles from it.
    """
    header = FIELDS["profile_url"].label
    rows, seen = [], set()
    for raw in urls:
        for piece in re.split(r"[\s,;]+", str(raw or "")):
            if not SLUG_RE.search(piece):
                continue
            slug = extract_slug(piece)
            if slug and slug not in seen:
                seen.add(slug)
                rows.append({header: profile_url(slug)})
    if not rows:
        raise InputError("None of those lines looked like a LinkedIn profile link.\n\n"
                         "They should look like https://www.linkedin.com/in/someone")
    return rows, [header]


# ===========================================================================
# section 13 -- writing the output
# ===========================================================================
def write_outputs(rows: list[dict], headers: list[str], folder: Path, stem: str,
                  want_xlsx: bool = True, want_csv: bool = True,
                  review_header: str = "Needs Review",
                  highlight: bool = True) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    if want_csv:
        path = folder / f"{stem}.csv"
        tmp = path.with_suffix(".csv.tmp")
        with tmp.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({h: row.get(h, "") for h in headers})
        _replace(tmp, path)
        written.append(path)

    if want_xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Candidates"
        sheet.append(headers)
        amber = PatternFill("solid", fgColor="FFF2CC")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            sheet.append([row.get(h, "") for h in headers])
            if highlight and str(row.get(review_header, "")).strip().lower() == "yes":
                for cell in sheet[sheet.max_row]:
                    cell.fill = amber
        sheet.freeze_panes = "A2"
        for i, header in enumerate(headers, start=1):
            width = max([len(str(header))] + [len(str(r.get(header, ""))) for r in rows]) \
                if rows else len(str(header))
            sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 12), 46)
        path = folder / f"{stem}.xlsx"
        tmp = path.with_suffix(".xlsx.tmp")
        workbook.save(tmp)
        _replace(tmp, path)
        written.append(path)

    return written


def _replace(tmp: Path, target: Path) -> None:
    """os.replace fails on Windows while the file is open in Excel -- say so clearly."""
    for attempt in range(3):
        try:
            os.replace(tmp, target)
            return
        except PermissionError:
            if attempt == 2:
                with contextlib.suppress(OSError):
                    tmp.unlink()
                raise OutputLockedError(
                    f"{target.name} is open in Excel, so it could not be saved.\n\n"
                    f"Please close it and press Export again.")
            time.sleep(0.5)


# ===========================================================================
# section 14 -- the runner
#
# A state machine the caller pumps, so a UI can report per-candidate progress,
# pause, and stop.  Everything from open_browser() onwards must happen on ONE
# thread: Playwright's synchronous API is thread-affine.
# ===========================================================================
@dataclass(frozen=True)
class Plan:
    total: int
    with_url: int
    need_search: int
    already_complete: int
    unreadable_links: tuple[str, ...]
    needs: frozenset
    headers: tuple[str, ...]
    est_pageviews: float
    est_seconds: float


@dataclass(frozen=True)
class Progress:
    index: int
    name: str
    stage: str


@dataclass(frozen=True)
class RowResult:
    index: int
    name: str
    slug: str | None
    status: str
    values: dict
    filled: tuple[str, ...]
    needs_review: bool
    confidence: str
    found_by: str
    notes: str
    from_cache: bool
    pageviews: int
    elapsed: float


@dataclass(frozen=True)
class RunSummary:
    reason: str
    counters: dict
    files: tuple[Path, ...]
    log_path: Path | None


class EnrichRunner:
    def __init__(self, config: AppConfig, control: Control | None = None,
                 creds: Credentials | None = None, log_path: Path | None = None):
        self.config = config
        self.control = control or Control()
        self.creds = creds
        self.log_path = log_path
        self.cache = RecordCache(config.cache_ttl_days, config.use_cache)

        self.rows: list[dict] = []
        self.headers: list[str] = []
        self._input_headers: list[str] = []
        self._link_header: str | None = None
        self._name_header: str | None = None
        self.plan: Plan | None = None
        self.counters = {"done": 0, "filled": 0, "review": 0, "not_found": 0,
                         "errors": 0, "skipped": 0, "pageviews": 0}
        self.results: dict[int, RowResult] = {}
        self.reason = "completed"
        self._session: Session | None = None
        self._pw_ctx = None
        self._pw = None
        self._durations: list[float] = []
        self._login_emit: Callable[[LoginEvent], None] = lambda ev: None

    # -- phase 1: no browser ----------------------------------------------
    def load(self) -> Plan:
        cfg = self.config
        problems = cfg.validate()
        if problems:
            raise MappingError("Please fix these settings first:\n\n"
                               + "\n".join(f"  • {p}" for p in problems))

        if cfg.input_mode == "urls":
            rows, headers = rows_from_urls(cfg.urls)
        else:
            if cfg.input_mode == "sheet":
                path = download_gsheet(cfg.sheet_url)
            else:
                path = Path(cfg.file_path).expanduser()
                if not path.exists():
                    raise InputError(f"That file no longer exists:\n\n{path}")
                if path.suffix.lower() not in (".csv", ".tsv", ".xlsx", ".xlsm"):
                    raise InputError(
                        f"'{path.suffix}' files are not supported. Please use an "
                        f"Excel (.xlsx) or CSV file.")
            rows, headers = load_rows(path)
        if not rows:
            raise InputError("That sheet has no data rows in it.")

        self._input_headers = list(headers)
        # output columns: the mapping's, then any input column we are leaving alone
        out_headers = [c.header for c in cfg.columns if c.enabled]
        passthrough = [h for h in headers if h not in out_headers]
        self.headers = out_headers + passthrough
        self.rows = rows
        if cfg.limit and cfg.limit > 0:
            self.rows = self.rows[: cfg.limit]

        link_header = self._input_header_for("profile_url")
        name_header = self._input_header_for("full_name")
        self._link_header, self._name_header = link_header, name_header
        needs = cfg.needs
        unreadable: list[str] = []

        for row in self.rows:
            raw_link = row.get(link_header, "") if link_header else ""
            slug = extract_slug(raw_link)
            if slug is None and not is_blank(raw_link):
                unreadable.append(str(raw_link))
            row["_slug"] = slug
            row["_missing"] = self._missing_fields(row)

        with_url = sum(1 for r in self.rows if r["_slug"])
        complete = sum(1 for r in self.rows if not r["_missing"])
        searchable = sum(1 for r in self.rows
                         if not r["_slug"] and not is_blank(r.get(name_header, ""))
                         and cfg.name_search)
        per_profile = self._estimate_pageviews(needs)
        work = with_url + searchable
        est = work * ((cfg.min_delay + cfg.max_delay) / 2 + per_profile * 4.0)

        self.plan = Plan(total=len(self.rows), with_url=with_url,
                         need_search=searchable, already_complete=complete,
                         unreadable_links=tuple(unreadable), needs=needs,
                         headers=tuple(self.headers), est_pageviews=per_profile,
                         est_seconds=est)
        log.info("plan: %s rows, %s with a link, %s to search, needs=%s",
                 len(self.rows), with_url, searchable, sorted(needs))
        return self.plan

    def _header_for(self, field_key: str) -> str | None:
        """Which output column carries this field, if any."""
        for column in self.config.columns:
            if column.field == field_key and column.enabled:
                return column.header
        return None

    def _input_header_for(self, field_key: str) -> str | None:
        """Which column to READ this field from.

        The LinkedIn link and the candidate name are inputs as much as outputs: a
        user may map only "Current Company" for export, but the sheet still has a
        LinkedIn column and we must read the links from it.  So check the mapping
        first, then fall back to recognising the input headers by name.
        """
        mapped = self._header_for(field_key)
        if mapped and mapped in self._input_headers:
            return mapped
        for header in self._input_headers:
            if LF.match_header_to_field(header) == field_key:
                return header
        return mapped

    def _missing_fields(self, row: dict) -> list[str]:
        missing = []
        for column in self.config.columns:
            if not column.enabled or not column.field:
                continue
            spec = FIELDS[column.field]
            if spec.overwrite == Overwrite.NEVER or not spec.needs:
                continue
            if is_blank(row.get(column.header)):
                missing.append(column.field)
        return missing

    def _page_gap(self) -> tuple[float, float]:
        """Pause between the extra pages of one profile: a fraction of the pacing
        the user chose, so it scales with their risk appetite instead of ignoring it."""
        low = max(0.0, self.config.min_delay * 0.2)
        high = max(low, self.config.max_delay * 0.25)
        return (low, high)

    def _estimate_pageviews(self, needs: frozenset) -> float:
        """Rough pages per profile: the main page, plus the extras that are needed.
        Education and experience only sometimes need their details page."""
        views = 1.0
        if LF.VISIT_EDUCATION in needs:
            views += 0.7
        if LF.VISIT_EXPERIENCE in needs:
            views += 0.45
        if LF.VISIT_CONTACT in needs:
            views += 1.0
        for visit in needs:
            if visit not in (LF.VISIT_PROFILE, LF.VISIT_EDUCATION,
                            LF.VISIT_EXPERIENCE, LF.VISIT_CONTACT):
                views += 0.6          # only the people who have that section
        return round(views, 2)

    # -- phase 2: browser -------------------------------------------------
    def open_browser(self) -> None:
        from playwright.sync_api import sync_playwright
        self._pw_ctx = sync_playwright()
        self._pw = self._pw_ctx.start()
        self._session = Session(self._pw, headless=self.config.headless,
                                control=self.control, creds=self.creds,
                                max_relogins=self.config.max_relogins)
        self._session.open()

    def session_state(self) -> str:
        if self._session is None:
            return "closed"
        return "live" if self._session.live() else "needs_login"

    def login(self, emit: Callable[[LoginEvent], None] | None = None) -> bool:
        if self._session is None:
            raise EngineError("The browser is not open yet.")
        self._login_emit = emit or (lambda ev: None)
        self._session.close()
        ok = assisted_login(self._pw, self.creds, self.control, self._login_emit)
        self._session.open()
        return bool(ok and self._session.live())

    # -- phase 3: the rows ------------------------------------------------
    def iter_rows(self) -> Iterator[RowResult | Progress]:
        """Yield a Progress for each sub-step and exactly one RowResult per row."""
        if self._session is None:
            raise EngineError("The browser is not open yet.")
        cfg = self.config
        name_header = self._name_header
        link_header = self._link_header
        processed = 0

        for index, row in enumerate(self.rows):
            name = str(row.get(name_header, "") or "").strip() if name_header else ""
            display = name or row.get(link_header, "") or f"row {index + 1}"
            started = time.monotonic()
            yield Progress(index, display, "start")

            try:
                result = self._process(index, row, display, name)
            except Cancelled:
                self.reason = "stopped"
                return
            except Exception as exc:
                log.error("row %s failed:\n%s", index + 1, traceback.format_exc())
                self.counters["errors"] += 1
                result = self._result(index, display, row.get("_slug"), "error", {},
                                      (), False, "", "", str(exc).splitlines()[0][:200],
                                      False, 0, time.monotonic() - started)

            self.results[index] = result
            self.counters["done"] += 1
            self.counters["pageviews"] += result.pageviews
            if result.pageviews:
                self._durations.append(result.elapsed)
            yield result

            if result.status == "blocked":
                self.reason = f"blocked:{result.notes or 'unknown'}"
                return

            if result.pageviews:
                processed += 1
            # pace only when we actually touched LinkedIn
            if result.pageviews and index < len(self.rows) - 1:
                try:
                    if processed and processed % max(1, cfg.long_pause_every) == 0:
                        pause = random.uniform(*cfg.long_pause)
                        log.info("long pause of %.0fs to stay under LinkedIn's limits", pause)
                        yield Progress(index, display, f"resting {pause:.0f}s")
                        self.control.sleep(pause)
                    else:
                        self.control.sleep(random.uniform(cfg.min_delay, cfg.max_delay))
                except Cancelled:
                    self.reason = "stopped"
                    return

    def _process(self, index: int, row: dict, display: str, name: str) -> RowResult:
        cfg = self.config
        started = time.monotonic()
        notes: list[str] = []
        slug = row.get("_slug")
        found_by = "sheet" if slug else ""
        confidence = ""
        needs_review = False

        if not row.get("_missing"):
            self.counters["skipped"] += 1
            return self._result(index, display, slug, "skipped", {}, (), False, "",
                                found_by, "nothing was missing", False, 0,
                                time.monotonic() - started)

        if not slug:
            if not cfg.name_search or not name:
                self.counters["skipped"] += 1
                return self._result(index, display, None, "skipped", {}, (), False, "",
                                    "", "no LinkedIn link to work from", False, 0,
                                    time.monotonic() - started)
            log.info("[%s] searching LinkedIn by name for %r", index + 1, name)
            email_header = self._input_header_for("email")
            email = str(row.get(email_header, "") or "") if email_header else ""
            found = resolve_missing_profile(self._session.page, name, email, self.control)
            if found["blocked"] and self._session.relogin(self._login_emit):
                found = resolve_missing_profile(self._session.page, name, email,
                                                self.control)
            if found["blocked"]:
                return self._result(index, display, None, "blocked", {}, (), True, "",
                                    "", found["blocked"], False, 1,
                                    time.monotonic() - started)
            if not found["slug"]:
                self.counters["not_found"] += 1
                return self._result(index, display, None, "not_found", {}, (), True, "",
                                    "", found["note"], False, 1,
                                    time.monotonic() - started)
            slug = found["slug"]
            found_by = "name search"
            confidence = found["confidence"]
            needs_review = confidence != "high"
            notes.append(f'matched "{found.get("matched_name", "")}" '
                         f'(score {found.get("score")})')
            if found["note"]:
                notes.append(found["note"])

        gap = self._page_gap()
        record = scrape_profile(self._session.page, slug, cfg.needs, self.control,
                               self.cache, gap)
        if record.get("blocked") and self._session.relogin(self._login_emit):
            record = scrape_profile(self._session.page, slug, cfg.needs, self.control,
                                   self.cache, gap)
        if record.get("blocked"):
            return self._result(index, display, slug, "blocked", {}, (), True,
                                confidence, found_by, record["blocked"], False,
                                record.get("pageviews", 1), time.monotonic() - started)

        ctx = Ctx(record, {})
        values = LF.derive_values(ctx, [c.field for c in cfg.columns
                                        if c.enabled and c.field])
        filled = self._merge(row, values)

        counts_edu = len(ctx.section_entries("education"))
        counts_exp = len(ctx.section_entries("experience"))
        if LF.VISIT_CONTACT in cfg.needs and not (values.get("email")
                                                  or values.get("phone")):
            notes.append("contact info not public")
        if LF.VISIT_EDUCATION in cfg.needs and not counts_edu:
            notes.append("no education listed on the profile")
        guessed = ctx.ug_pg[2]
        if guessed and any(f"{slot}_college" in filled for slot in guessed):
            notes.append("/".join(sorted(guessed)).upper()
                         + " guessed from dates -- the profile does not name the degree")
            needs_review = True
        full_list = values.get("all_companies_full") or ctx.companies[1]
        if values.get("all_companies") and full_list != values.get("all_companies"):
            notes.append("all employers: " + full_list)

        status = "ok" if filled else "nothing_new"
        if filled:
            self.counters["filled"] += 1
        if needs_review:
            self.counters["review"] += 1

        # meta columns are written last, from the row's own outcome
        meta = {"status": status, "notes": "; ".join(n for n in notes if n),
                "needs_review": "yes" if needs_review else "",
                "confidence": confidence, "found_by": found_by,
                "scraped_at": datetime.now(timezone.utc).astimezone()
                                      .strftime("%Y-%m-%d %H:%M"),
                "pageviews": record.get("pageviews", 0),
                "row_number": index + 1}
        meta_values = LF.derive_values(Ctx(record, meta),
                                       [c.field for c in cfg.columns
                                        if c.enabled and c.field
                                        and FIELDS[c.field].group == "Meta"])
        self._merge(row, meta_values)
        values.update(meta_values)

        log.info("%s -> filled %s (edu=%s exp=%s, %s pages)", slug, list(filled),
                 counts_edu, counts_exp, record.get("pageviews", 0))
        return self._result(index, display, slug, status, self._row_values(row), filled,
                            needs_review, confidence, found_by, meta["notes"],
                            bool(record.get("from_cache")), record.get("pageviews", 0),
                            time.monotonic() - started)

    def _merge(self, row: dict, values: dict) -> tuple[str, ...]:
        """Write values into the row, honouring each field's overwrite policy."""
        filled = []
        for column in self.config.columns:
            if not column.enabled or not column.field:
                continue
            spec = FIELDS[column.field]
            value = values.get(column.field)
            if value is None or is_blank(value):
                continue
            if spec.overwrite == Overwrite.NEVER:
                continue
            if spec.overwrite == Overwrite.IF_BLANK and not is_blank(row.get(column.header)):
                continue
            row[column.header] = str(value).strip()
            filled.append(column.field)
        return tuple(filled)

    def _row_values(self, row: dict) -> dict:
        return {h: row.get(h, "") for h in self.headers}

    def _result(self, index, name, slug, status, values, filled, review, confidence,
                found_by, notes, from_cache, pageviews, elapsed) -> RowResult:
        row = self.rows[index]
        return RowResult(index=index, name=name, slug=slug, status=status,
                         values=values or self._row_values(row), filled=tuple(filled),
                         needs_review=review, confidence=confidence, found_by=found_by,
                         notes=notes, from_cache=from_cache, pageviews=pageviews,
                         elapsed=elapsed)

    # -- estimates and finishing ------------------------------------------
    def eta_seconds(self) -> float:
        remaining = sum(1 for i, r in enumerate(self.rows)
                        if i not in self.results and r.get("_missing"))
        if not remaining:
            return 0.0
        if len(self._durations) < 3:
            return -1.0
        recent = sorted(self._durations[-10:])
        median = recent[len(recent) // 2]
        pacing = (self.config.min_delay + self.config.max_delay) / 2
        return remaining * (median + pacing)

    def export_rows(self) -> tuple[list[dict], list[str]]:
        rows = []
        for index, row in enumerate(self.rows):
            values = dict(self._row_values(row))
            rows.append(values)
        return rows, list(self.headers)

    def finish(self) -> RunSummary:
        cfg = self.config
        rows, headers = self.export_rows()
        review_header = self._header_for("needs_review") or "Needs Review"
        files: list[Path] = []
        try:
            files = write_outputs(rows, headers, Path(cfg.output_folder).expanduser(),
                                  cfg.output_stem, cfg.output_xlsx, cfg.output_csv,
                                  review_header, cfg.highlight_needs_review)
        except OutputLockedError:
            raise
        except Exception as exc:
            log.error("could not write the output: %s", exc)
            raise EngineError(f"Could not save the results: {exc}") from exc
        return RunSummary(reason=self.reason, counters=dict(self.counters),
                          files=tuple(files), log_path=self.log_path)

    def rederive(self) -> list[RowResult]:
        """Re-apply the current mapping to cached records, with no browser at all.

        This is what makes raw-record caching worth it: renaming or adding a column
        costs nothing and shows results instantly.
        """
        out = []
        for index, row in enumerate(self.rows):
            slug = row.get("_slug")
            if not slug:
                continue
            record, missing = self.cache.get(slug, self.config.needs)
            if record is None:
                continue
            ctx = Ctx(record, {"status": "from cache", "row_number": index + 1})
            values = LF.derive_values(ctx, [c.field for c in self.config.columns
                                            if c.enabled and c.field])
            filled = self._merge(row, values)
            out.append(self._result(index, row.get("_slug") or "", slug,
                                    "cached" if not missing else "partial",
                                    self._row_values(row), filled, False, "", "sheet",
                                    "" if not missing else "some pages not cached yet",
                                    True, 0, 0.0))
        return out

    def close(self) -> None:
        if self._session is not None:
            self._session.close()
            self._session = None
        if self._pw_ctx is not None:
            with contextlib.suppress(Exception):
                self._pw_ctx.__exit__(None, None, None)
            self._pw_ctx, self._pw = None, None
